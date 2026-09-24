import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook

from gridtoev.eirgrid_history import DataQualityError
from gridtoev.outage_constraints import (
    build_outage_ablation_report,
    build_outage_constraint_features,
    parse_ecp_constraint_report,
    parse_generation_outage_plan,
    parse_transmission_outage_programme,
    validate_outage_feature_matrix,
)


def outage(
    *,
    outage_id: str,
    family: str,
    published: str,
    start: str,
    end: str,
    outage_type: str,
    unavailable_mw: float | None = None,
    region: str = "ireland",
    voltage_kv: float | None = None,
) -> dict[str, object]:
    return {
        "outage_id": outage_id,
        "source_family": family,
        "snapshot_id": f"{family}:{published}",
        "published_at_utc": published,
        "start_timestamp_utc": start,
        "end_timestamp_utc": end,
        "outage_type": outage_type,
        "asset_name": outage_id,
        "unit_id": outage_id,
        "region": region,
        "constraint_group": region,
        "rated_mw": unavailable_mw,
        "unavailable_mw": unavailable_mw,
        "voltage_kv": voltage_kv,
        "source_url": "https://example.test/source.xlsx",
        "source_file": "source.xlsx",
        "schema_version": "eirgrid-outages-v1",
    }


class OutageConstraintParserTests(unittest.TestCase):
    def test_generation_parser_separates_forced_scheduled_and_derating(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "generation.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Outage Dates"
            sheet.append([None, "All Island Outage Start and End Times"])
            sheet.append([])
            sheet.append([])
            sheet.append(
                [None, "Station", "Unit", "NPR", "Start Date", "End Date", "Value", "Duration"]
            )
            sheet.append(
                [None, "Aghada", "AD1", 100, "2026-01-05 00:00", "2026-01-06 00:00", "F", 1]
            )
            sheet.append(
                [None, "Kilroot", "NIK1", 80, "2026-01-07 00:00", "2026-01-08 00:00", "S", 1]
            )
            sheet.append(
                [None, "Tarbert", "TB1", 160, "2026-01-09 00:00", "2026-01-10 00:00", 50, 1]
            )
            workbook.save(path)

            parsed = parse_generation_outage_plan(
                path,
                published_at="2026-01-01T12:00:00Z",
                source_url="https://example.test/generation.xlsx",
            )

        self.assertEqual(list(parsed["outage_type"]), ["forced", "planned", "derating"])
        self.assertEqual(list(parsed["unavailable_mw"]), [100.0, 80.0, 110.0])
        self.assertEqual(list(parsed["region"]), ["ireland", "northern_ireland", "ireland"])
        self.assertTrue((parsed["published_at_utc"] == pd.Timestamp("2026-01-01T12:00:00Z")).all())

    def test_transmission_and_ecp_parsers_normalise_regions_units_and_pressure(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            transmission_path = Path(directory) / "transmission.xlsx"
            transmission = Workbook()
            all_sheet = transmission.active
            all_sheet.title = "GEN_ALL"
            all_sheet.append(
                [
                    "Outage ID",
                    "Feeder ID",
                    "Calendar Day Duration",
                    "Working Day Duration",
                    "Work Days",
                    "Outage Status",
                    "Start",
                    "Finish",
                ]
            )
            all_sheet.append(
                ["TO-1", "220kV FEEDER - A 220-B 220-1", 2, 2, 2, "Scheduled", "05/01/2026", "06/01/2026"]
            )
            regional = transmission.create_sheet("GEN_Dublin")
            regional.append(list(next(all_sheet.values)))
            regional.append(list(next(iter(all_sheet.iter_rows(min_row=2, values_only=True)))))
            transmission.save(transmission_path)

            ecp_path = Path(directory) / "ecp.xlsx"
            ecp = Workbook()
            info = ecp.active
            info.title = "Info"
            results = ecp.create_sheet("Dispatch Down Results")
            results.append([None])
            results.append(
                [
                    "Area",
                    "Node",
                    "Year",
                    "Generation Scenario",
                    "Generation Type",
                    "Installed (MW)",
                    "Available Energy (GWh)",
                    "Surplus (GWh)",
                    "Curtailment (GWh)",
                    "Constraint (GWh)",
                    "Total Dispatch Down (GWh)",
                    "Surplus (%)",
                    "Curtailment (%)",
                    "Constraint (%)",
                ]
            )
            results.append(["A", "node-1", 2028, "Initial", "wind", 100, 300, 0, 0, 30, 30, 0, 0, 0.1])
            ecp.save(ecp_path)

            parsed_transmission = parse_transmission_outage_programme(
                transmission_path,
                published_at="2026-01-02T00:00:00Z",
                source_url="https://example.test/transmission.xlsx",
            )
            parsed_ecp = parse_ecp_constraint_report(
                ecp_path,
                published_at="2026-03-01T00:00:00Z",
                source_url="https://example.test/ecp.xlsx",
            )

        self.assertEqual(len(parsed_transmission), 1)
        self.assertEqual(parsed_transmission.loc[0, "region"], "dublin")
        self.assertEqual(parsed_transmission.loc[0, "voltage_kv"], 220.0)
        self.assertEqual(
            parsed_transmission.loc[0, "end_timestamp_utc"],
            pd.Timestamp("2026-01-07T00:00:00Z"),
        )
        self.assertEqual(len(parsed_ecp), 1)
        self.assertEqual(parsed_ecp.loc[0, "constraint_group"], "A")
        self.assertAlmostEqual(parsed_ecp.loc[0, "constraint_ratio"], 0.1)


class OutageConstraintFeatureTests(unittest.TestCase):
    def test_overlaps_boundaries_and_publication_vintages_are_causal(self) -> None:
        modelling = pd.DataFrame(
            {
                "issue_timestamp_utc": [
                    "2026-01-01T10:00:00Z",
                    "2026-01-01T10:30:00Z",
                    "2026-01-01T12:00:00Z",
                ],
                "target_timestamp_utc": [
                    "2026-01-01T10:30:00Z",
                    "2026-01-01T11:00:00Z",
                    "2026-01-01T12:30:00Z",
                ],
                "forecast_horizon_minutes": [30, 30, 30],
            }
        )
        outages = pd.DataFrame(
            [
                outage(
                    outage_id="forced-old",
                    family="generation_plan",
                    published="2026-01-01T08:00:00Z",
                    start="2026-01-01T10:00:00Z",
                    end="2026-01-01T13:00:00Z",
                    outage_type="forced",
                    unavailable_mw=100,
                ),
                outage(
                    outage_id="planned-boundary",
                    family="generation_plan",
                    published="2026-01-01T08:00:00Z",
                    start="2026-01-01T10:30:00Z",
                    end="2026-01-01T11:00:00Z",
                    outage_type="planned",
                    unavailable_mw=50,
                ),
                outage(
                    outage_id="forced-revised",
                    family="generation_plan",
                    published="2026-01-01T11:00:00Z",
                    start="2026-01-01T10:00:00Z",
                    end="2026-01-01T13:00:00Z",
                    outage_type="forced",
                    unavailable_mw=300,
                ),
                outage(
                    outage_id="network-1",
                    family="transmission_programme",
                    published="2026-01-01T08:00:00Z",
                    start="2026-01-01T09:00:00Z",
                    end="2026-01-01T12:00:00Z",
                    outage_type="planned",
                    region="dublin",
                    voltage_kv=220,
                ),
            ]
        )
        constraints = pd.DataFrame(
            {
                "constraint_group": ["A", "A"],
                "node": ["one", "one"],
                "study_year": [2028, 2028],
                "scenario": ["Initial", "Initial"],
                "generation_type": ["wind", "wind"],
                "installed_mw": [100.0, 100.0],
                "available_energy_gwh": [300.0, 300.0],
                "constraint_gwh": [60.0, 240.0],
                "constraint_ratio": [0.2, 0.8],
                "snapshot_id": ["ecp-old", "ecp-new"],
                "published_at_utc": ["2026-01-01T08:00:00Z", "2026-01-01T11:00:00Z"],
                "source_url": ["https://example.test/ecp.xlsx"] * 2,
                "source_file": ["ecp.xlsx"] * 2,
                "schema_version": ["eirgrid-ecp-v1"] * 2,
            }
        )

        features, dictionary, quality = build_outage_constraint_features(
            modelling, outages, constraints
        )

        first = features.iloc[0]
        at_end = features.iloc[1]
        after_revision = features.iloc[2]
        self.assertEqual(first["generation_active_outage_count"], 2)
        self.assertEqual(first["generation_forced_unavailable_mw"], 100)
        self.assertEqual(first["generation_planned_unavailable_mw"], 50)
        self.assertEqual(first["generation_largest_unavailable_unit_mw"], 100)
        self.assertEqual(at_end["generation_active_outage_count"], 1)
        self.assertEqual(at_end["generation_planned_unavailable_mw"], 0)
        self.assertEqual(after_revision["generation_forced_unavailable_mw"], 300)
        self.assertAlmostEqual(first["ecp_constraint_pressure_weighted_ratio"], 0.2)
        self.assertAlmostEqual(after_revision["ecp_constraint_pressure_weighted_ratio"], 0.8)
        self.assertEqual(first["transmission_active_outage_count"], 1)
        self.assertEqual(at_end["transmission_active_outage_count"], 1)
        self.assertEqual(after_revision["transmission_active_outage_count"], 0)
        self.assertEqual(len(features), len(modelling))
        self.assertEqual(set(features.columns), set(dictionary["column"]))
        self.assertEqual(quality["future_publication_violations"], 0)
        self.assertEqual(quality["duplicate_natural_keys"], 0)

    def test_validator_rejects_future_publication_and_duplicate_rows(self) -> None:
        table = pd.DataFrame(
            {
                "issue_timestamp_utc": ["2026-01-01T10:00:00Z"] * 2,
                "target_timestamp_utc": ["2026-01-01T10:30:00Z"] * 2,
                "forecast_horizon_minutes": [30, 30],
                "generation_snapshot_published_at_utc": ["2026-01-01T11:00:00Z"] * 2,
            }
        )
        with self.assertRaises(DataQualityError):
            validate_outage_feature_matrix(table)

    def test_ablation_uses_development_folds_and_never_scores_final_test(self) -> None:
        issue = pd.date_range("2026-01-01", periods=200, freq="30min", tz="UTC")
        signal = np.tile([0.0, 100.0], 100)
        model_data = pd.DataFrame(
            {
                "issue_timestamp_utc": issue,
                "target_timestamp_utc": issue + pd.Timedelta(minutes=30),
                "forecast_horizon_minutes": 30,
                "dispatch_down_mwh_latest_observed": 0.0,
                "baseline_noise": np.sin(np.arange(200)),
                "dispatch_down_mwh": signal,
            }
        )
        feature_table = model_data[
            [
                "issue_timestamp_utc",
                "target_timestamp_utc",
                "forecast_horizon_minutes",
            ]
        ].copy()
        feature_table["generation_total_unavailable_mw"] = signal

        with tempfile.TemporaryDirectory() as directory:
            artifact = Path(directory) / "features.csv"
            feature_table.to_csv(artifact, index=False)
            report = build_outage_ablation_report(
                model_data,
                feature_table,
                artifact,
            )

        self.assertEqual(report["experiment_name"], "outage_constraint_signals")
        self.assertEqual(report["development_fold_count"], 4)
        self.assertGreater(report["metrics"]["mean_mae_improvement_fraction"], 0.5)
        self.assertEqual(report["production_decision"], "include")
        self.assertFalse(report["final_test_accessed"])


if __name__ == "__main__":
    unittest.main()
